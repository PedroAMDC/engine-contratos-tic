"""Mapa preliminar da infraestrutura de dados do Governo, a partir dos contratos.

Responde a levantamentos sobre onde estão os dados de cada
órgão e quanto custa manter a infraestrutura. Complementa o resumo estadual (tema 1) e lista,
por órgão, o que ainda precisa ser informado pelas equipes.
"""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import date
from pathlib import Path

from ..config import DATA, OUTPUT
from ..perfil import P
from .base import GLOBAL, carregar_base

DC_ESTATAL = f"Datacenter {P.estatal('de')}"
TIPOS = [DC_ESTATAL, "Outro datacenter governamental", "Nuvem pública", "Nuvem privada de terceiros",
         "Colocation", "Ambiente próprio do órgão", "Não identificado"]
SUBS = {"Link de internet e rede de dados", "Rede local e Wi-Fi", "Segurança da informação (software)",
        "Segurança gerenciada", "Banco de dados e infraestrutura de software"}
TEMAS = {"nuvem", "hospedagem", "data center", "backup", "firewall/waf", "banco de dados"} | (
    {P.estatal_ti.rede_governo.lower()} if P.estatal_ti.rede_governo else set())
LOTE = 8


def selecionar() -> list[dict]:
    base = carregar_base()
    alvo = [c for c in base if c["macro"] == "Infraestrutura e nuvem" or c["sub"] in SUBS or c["estatal_ti"]
            or set(c["temas"]) & TEMAS or c.get("nuvem")]
    por_pasta = defaultdict(list)
    for c in alvo:
        por_pasta[c["id"].rsplit("-", 1)[0]].append(c["id"])
    lotes = []
    for pasta, ids in sorted(por_pasta.items()):
        for i in range(0, len(ids), LOTE):
            lotes.append({"id": f"INFRA-{pasta}-{i // LOTE + 1:02d}", "pasta": pasta, "contratos": ids[i:i + LOTE]})
    (GLOBAL / "lotes_infra.json").write_text(json.dumps(lotes, ensure_ascii=False, indent=1), encoding="utf-8")
    return lotes


def verificar_infra(pasta: str, grupo: str) -> dict:
    from ..verificacao import CPF, _num, conferir_fonte

    arq = DATA / pasta / "infra" / f"{grupo}.json"
    erros, avisos = [], []
    try:
        x = json.loads(arq.read_text(encoding="utf-8"))
    except Exception as e:
        return {"ok": False, "erros": [f"JSON inválido ou ausente: {e}"], "avisos": []}
    ficha = json.loads((DATA / pasta / "fichas" / f"{grupo}.json").read_text(encoding="utf-8"))
    docs = set(ficha.get("documentos", []))
    for k in ("pagador", "relevante", "ambientes", "lacunas"):
        if k not in x:
            erros.append(f"campo obrigatório ausente: {k}")
    for i, a in enumerate(x.get("ambientes", [])):
        if a.get("tipo") not in TIPOS:
            erros.append(f"ambientes[{i}].tipo fora da lista: {TIPOS}")
        for k in ("custo_mensal", "custo_anual"):
            if a.get(k) is not None and not _num(a[k]):
                erros.append(f"ambientes[{i}].{k} deve ser número ou null")
        if a.get("tipo") != "Não identificado":
            if not a.get("fonte"):
                erros.append(f"ambientes[{i}] sem fonte")
            else:
                conferir_fonte(pasta, f"ambientes[{i}]", a["fonte"], docs, True, erros, avisos)
    if CPF.search(json.dumps(x, ensure_ascii=False)):
        erros.append("contém CPF")
    arq.write_text(json.dumps(x, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"ok": not erros, "erros": erros, "avisos": avisos}


def montar_mapa() -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    base = {c["id"]: c for c in carregar_base()}
    regs = []
    for arq in sorted(DATA.glob("*/infra/G*.json")):
        pasta = arq.parent.parent.name
        x = json.loads(arq.read_text(encoding="utf-8"))
        x["_id"] = f"{pasta}-{arq.stem}"
        regs.append(x)

    orgaos = sorted({c["orgao"] for c in base.values()})
    por_orgao = {o: {"ambientes": defaultdict(float), "tipos": set(), "contratos": 0, "custo_anual": 0.0,
                     "sem_custo": 0, "lacunas": set()} for o in orgaos}
    linhas_amb = []
    for x in regs:
        c = base.get(x["_id"])
        if not c:
            continue
        org = (x.get("pagador") or c["orgao"]).upper()
        d = por_orgao.setdefault(org, {"ambientes": defaultdict(float), "tipos": set(), "contratos": 0,
                                       "custo_anual": 0.0, "sem_custo": 0, "lacunas": set()})
        if x.get("relevante"):
            d["contratos"] += 1
        d["lacunas"].update(x.get("lacunas", []))
        for a in x.get("ambientes", []):
            d["tipos"].add(a["tipo"])
            anual = a.get("custo_anual") or ((a.get("custo_mensal") or 0) * 12) or None
            if anual:
                d["custo_anual"] += anual
            else:
                d["sem_custo"] += 1
            linhas_amb.append([org, c["id"], c["contratada"], a["tipo"], a.get("provedor"), a.get("local"), a.get("endereco"),
                               ", ".join(a.get("componentes", [])), a.get("custo_mensal"), anual, a.get("base_custo"),
                               f"{a['fonte']['doc']} p.{a['fonte']['pagina']}" if a.get("fonte") else None, c["situacao"]])

    destino = OUTPUT / "ESTADUAL" / f"Mapa_Infraestrutura_Dados_{P.ente.sigla}_preliminar.xlsx"
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    cab = PatternFill("solid", fgColor="0F2A4A")

    def aba(titulo, colunas, linhas, moeda=()):
        ws = wb.create_sheet(titulo)
        ws.append(colunas)
        for cel in ws[1]:
            cel.font, cel.fill, cel.alignment = Font(bold=True, color="FFFFFF"), cab, Alignment(wrap_text=True)
        for l in linhas:
            ws.append(l)
        for i, col in enumerate(colunas, 1):
            letra = get_column_letter(i)
            ws.column_dimensions[letra].width = min(55, max(12, len(col) + 2))
            if col in moeda:
                for cel in ws[letra][1:]:
                    cel.number_format = '"R$" #,##0.00'
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb.remove(wb.active)
    resumo = []
    for org, d in sorted(por_orgao.items()):
        tipos = sorted(d["tipos"] - {"Não identificado"})
        pedir = ["Endereço ou georreferenciamento dos ambientes"]
        if not tipos:
            pedir.insert(0, "Onde estão os dados (nenhum contrato de infraestrutura de dados nos documentos: estrutura "
                              f"{P.estatal('de')} sem contrato próprio, ambiente próprio ou terceiros?)")
        if "Ambiente próprio do órgão" in tipos or not tipos:
            pedir.append("Custos estimados da infraestrutura própria (energia, climatização, cabeamento, equipe)")
        if DC_ESTATAL in tipos:
            pedir.append(f"Confirmar se o valor pago {P.estatal('a')} cobre todo o ambiente ou se há rateio adicional")
        if d["sem_custo"]:
            pedir.append(f"Custos de {d['sem_custo']} ambiente(s) sem valor identificado no contrato")
        resumo.append([org, ", ".join(tipos) or "Não identificado nos contratos", d["contratos"], d["custo_anual"] or None,
                       "; ".join(sorted(d["lacunas"]))[:500] or None, " | ".join(pedir)])
    aba("Por órgão", ["Órgão", "Ambientes identificados", "Contratos de infraestrutura", "Custo anual identificado",
                      "Lacunas apontadas nos contratos", "O que pedir ao órgão"], resumo, moeda=("Custo anual identificado",))
    aba("Ambientes", ["Órgão pagador", "Contrato", "Contratada", "Tipo de ambiente", "Provedor", "Local", "Endereço",
                      "Componentes", "Custo mensal", "Custo anual", "Base do custo", "Fonte", "Situação"],
        linhas_amb, moeda=("Custo mensal", "Custo anual"))
    ws = wb.create_sheet("Sobre", 0)
    for l in [f"Mapa preliminar da infraestrutura de dados do {P.ente.governo} — a partir dos contratos",
              f"Gerado em {date.today():%d/%m/%Y} {P.emissor_com('por')} — {P.emissor_vinculo}.",
              "Fonte: contratos de TIC enviados pelos órgãos. Não inclui infraestrutura própria sem contrato, energia, climatização "
              f"ou cabeamento contratados fora da área de TIC, nem rateios internos {P.estatal('de')}.",
              "Aba 'Por órgão', coluna 'O que pedir ao órgão': informações que as equipes "
              f"{P.central('de')} e {P.estatal('de')} devem solicitar.",
              "Custos: valores do que cada contrato paga pelo ambiente, quando documentados ou calculáveis."]:
        ws.append([l])
    ws["A1"].font = Font(bold=True, size=13, color="0F2A4A")
    ws.column_dimensions["A"].width = 130
    wb.save(destino)
    return destino


def resumo_mapa() -> dict | None:
    """Resumo do mapa de infraestrutura para o resumo executivo estadual."""
    base = {c["id"]: c for c in carregar_base()}
    regs = []
    for arq in sorted(DATA.glob("*/infra/G*.json")):
        x = json.loads(arq.read_text(encoding="utf-8"))
        x["_id"] = f"{arq.parent.parent.name}-{arq.stem}"
        if x["_id"] in base:
            regs.append(x)
    if not regs:
        return None
    por_tipo = {t: {"ambientes": 0, "custo_anual": 0.0, "orgaos": set()} for t in TIPOS}
    orgaos_tipos = defaultdict(set)
    for x in regs:
        org = (x.get("pagador") or base[x["_id"]]["orgao"]).upper()
        for a in x.get("ambientes", []):
            d = por_tipo[a["tipo"]]
            d["ambientes"] += 1
            d["custo_anual"] += a.get("custo_anual") or ((a.get("custo_mensal") or 0) * 12)
            d["orgaos"].add(org)
            orgaos_tipos[org].add(a["tipo"])
    todos = {c["orgao"] for c in base.values()}
    identificados = {o for o, ts in orgaos_tipos.items() if ts - {"Não identificado"}}
    return {
        "por_tipo": [{"tipo": t, "ambientes": v["ambientes"], "custo_anual": round(v["custo_anual"], 2), "orgaos": len(v["orgaos"])}
                     for t, v in por_tipo.items() if v["ambientes"]],
        "custo_total": round(sum(v["custo_anual"] for v in por_tipo.values()), 2),
        "ambientes": sum(v["ambientes"] for v in por_tipo.values()),
        "multiplos": sorted(o for o, ts in orgaos_tipos.items() if len(ts - {"Não identificado"}) >= 2),
        "so_estatal_ti": sorted(o for o, ts in orgaos_tipos.items() if ts - {"Não identificado"} == {DC_ESTATAL}),
        "proprios": sorted(o for o, ts in orgaos_tipos.items() if "Ambiente próprio do órgão" in ts),
        "sem_ambiente": sorted(todos - identificados),
        "com_endereco": sum(1 for x in regs for a in x.get("ambientes", []) if a.get("endereco")),
    }
